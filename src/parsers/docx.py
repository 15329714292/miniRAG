import os

import docx

from .base import BaseParser, Document
from ..utils.text import normalize_text


class DocxParser(BaseParser):
    def parse(self, file_path: str):
        file_name = os.path.basename(file_path)
        # Load document and join non-empty paragraphs.
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        text = normalize_text("\n".join(paragraphs))
        if not text:
            return []
        # Build a single document record for the file.
        return [Document(file_path=file_path, file_name=file_name, text=text)]
